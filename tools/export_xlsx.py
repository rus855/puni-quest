# 問題の 一覧（tools/questions.json）から Excel を つくる
# つかいかた：python tools/export_xlsx.py [出力ファイル名] [国のID（sci など。つけると その国だけ）]
import json, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

here = os.path.dirname(os.path.abspath(__file__))
data = json.load(open(os.path.join(here, "questions.json"), encoding="utf-8"))
out_dir = os.path.join(os.path.dirname(os.path.dirname(here)), "ぷにまる問題一覧")
os.makedirs(out_dir, exist_ok=True)
name = sys.argv[1] if len(sys.argv) > 1 else f"ぷにまるクエスト_問題一覧_v{data['v']}_{datetime.date.today():%Y%m%d}.xlsx"
path = os.path.join(out_dir, name)
only = sys.argv[2] if len(sys.argv) > 2 else None
if only:
    data["worlds"] = [w for w in data["worlds"] if w["id"] == only]
    data["chapters"] = [c for c in data["chapters"] if c["w"] == only]
    data["rows"] = [r for r in data["rows"] if r["world"] == only]
    if only != "kanji":
        data["extra"] = []

wb = Workbook()
head_font = Font(bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="E0528E")
wrap = Alignment(wrap_text=True, vertical="top")

# まとめの シート
ws = wb.active
ws.title = "まとめ"
ws.append(["国", "章ID", "章の名前", "学年", "中学受験用", "出題項目の数", "一覧の行数"])
cnt = {}
for r in data["rows"]:
    cnt[r["ch"]] = cnt.get(r["ch"], 0) + 1
wname = {w["id"]: w["name"] for w in data["worlds"]}
for c in data["chapters"]:
    ws.append([wname[c["w"]], c["id"], c["name"], c["g"], "中学受験用" if c["exam"] else "", c["items"], cnt.get(c["id"], 0)])
for e in data.get("extra", []):
    ws.append(["漢字の国", e["id"], e["name"] + f"（{e['g']}年）", e["g"], "", e["n"], "（その学年の漢字を書く）"])
for i, w in enumerate([10, 8, 28, 6, 10, 12, 10], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for cell in ws[1]:
    cell.font = head_font; cell.fill = head_fill

# 国ごとの シート
cols = ["章ID", "章の名前", "学年", "中学受験用", "出はじめる ステージ", "項目ID", "カード名", "聞き方", "問題", "正解", "まちがい1", "まちがい2", "まちがい3", "解説"]
widths = [7, 22, 5, 9, 9, 12, 16, 14, 46, 22, 20, 20, 20, 36]
for w in data["worlds"]:
    sh = wb.create_sheet(w["name"])
    sh.append(cols)
    for r in data["rows"]:
        if r["world"] != w["id"]:
            continue
        ng = (r.get("ng") or []) + ["", "", ""]
        sh.append([r["ch"], r["chName"], r["g"], r["exam"], r["stage"], r["id"], r.get("card", ""), r["kind"], r["q"], str(r["ans"]), str(ng[0]), str(ng[1]), str(ng[2]), r.get("ex", "")])
    for i, wd in enumerate(widths, 1):
        sh.column_dimensions[get_column_letter(i)].width = wd
    for cell in sh[1]:
        cell.font = head_font; cell.fill = head_fill
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = sh.dimensions

# せつめいの シート
info = wb.create_sheet("見かた")
for line in [
    "この一覧は、ゲームの中の問題データから自動で作っています。",
    "・「出はじめる ステージ」：章の問題は、ステージ1では前の1/3（少ない章は前の10種類）、ステージ2ではその2倍、ステージ3とボスでは全部から出ます。",
    "・漢字の国：1つの漢字につき、熟語・ことばごとに「読み」と「書き」を載せています（ゲームでは、ほかに「漢字で書くと？」「共通の漢字」の聞き方もあります）。",
    "・漢字の「読み」のまちがいの選択肢は、出題のたびに変わります。ここに載せているのは、その一例です。",
    "・計算（例）・公式：計算（例）は、出題のたびに数字が変わります。ここには例を数問載せています。",
    "・「ぎゃくに きく」は、答えから問題を当てる聞き方です。まちがいの選択肢は、同じ章のほかの答えから選ばれます。",
    "・ゲームでは、ほかに「この答えは正しい？（⭕❌）」の聞き方もあります。",
]:
    info.append([line])
info.column_dimensions["A"].width = 120
wb.move_sheet("見かた", offset=-(len(wb.sheetnames) - 1))
wb.save(path)
print("saved", path)
